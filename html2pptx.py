from bs4 import BeautifulSoup
from pptx import Presentation
from pptx.util import Inches, Pt
import sys
import re

html_file = r"c:\Users\ADMIN\Desktop\Aarogya Sathi Capstone\Aarogya_Sathi_Presentation.html"
pptx_file = r"c:\Users\ADMIN\Desktop\Aarogya Sathi Capstone\Aarogya_Sathi_Presentation.pptx"

def extract_text(element):
    if not element:
        return ""
    return re.sub(r'\s+', ' ', element.get_text()).strip()

def create_presentation():
    try:
        with open(html_file, 'r', encoding='utf-8') as f:
            html = f.read()
    except Exception as e:
        print(f"Error reading {html_file}: {e}")
        return

    soup = BeautifulSoup(html, 'html.parser')
    prs = Presentation()

    slides = soup.find_all('div', class_='slide')

    for slide_div in slides:
        # Check if it's the cover slide
        if 'cover-slide' in slide_div.get('class', []):
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            title = slide.shapes.title
            subtitle = slide.placeholders[1]
            
            h1 = slide_div.find('h1', class_='cover-title')
            sub = slide_div.find('div', class_='cover-sub')
            
            title.text = extract_text(h1) if h1 else "Aarogya Sathi"
            subtitle.text = extract_text(sub) if sub else "AI-Powered Health Assistant"
            continue

        # Standard slide
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        title = slide.shapes.title
        content = slide.placeholders[1]
        
        # Extract title
        h2 = slide_div.find('h2')
        header = slide_div.find('h4')
        
        title_text = ""
        if header:
            title_text += extract_text(header) + "\n"
        if h2:
            title_text += extract_text(h2)
            
        title.text = title_text

        # Extract content bodies
        tf = content.text_frame
        tf.clear()  # clear default paragraph
        tf.word_wrap = True

        def add_paragraph(text, level=0, bold=False):
            if not text.strip(): return
            p = tf.add_paragraph()
            p.text = text
            p.level = level
            if bold:
                p.font.bold = True

        # Remove header and title so we don't extract them again
        if header: header.decompose()
        if h2: h2.decompose()
        
        # To prevent duplicate nested text, we extract structural blocks
        blocks = slide_div.find_all(['h3', 'p', 'li', 'blockquote', 'div'])
        
        processed_tags = set()
        
        for tag in blocks:
            # Skip if any parent has already been processed (prevents nested duplicates)
            parent = tag.parent
            skip = False
            while parent:
                if parent in processed_tags:
                    skip = True
                    break
                parent = parent.parent
            if skip:
                continue

            if tag.name == 'div':
                classes = tag.get('class', [])
                if 'stat-val' in classes or 'stat-label' in classes or 'tech-name' in classes or 'outcome-metric' in classes:
                    add_paragraph(extract_text(tag), level=0, bold=True)
                    processed_tags.add(tag)
                elif 'stat-sub' in classes or 'tech-desc' in classes:
                    add_paragraph(extract_text(tag), level=1)
                    processed_tags.add(tag)
            elif tag.name == 'h3':
                add_paragraph(extract_text(tag), level=0, bold=True)
                processed_tags.add(tag)
            elif tag.name in ['p', 'blockquote']:
                add_paragraph(extract_text(tag), level=1)
                processed_tags.add(tag)
            elif tag.name == 'li':
                add_paragraph(extract_text(tag), level=1)
                processed_tags.add(tag)

    prs.save(pptx_file)
    print(f"Successfully saved to {pptx_file}")

if __name__ == '__main__':
    create_presentation()
